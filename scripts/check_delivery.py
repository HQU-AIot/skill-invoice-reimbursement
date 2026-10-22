#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Final audit of a reimbursement delivery folder.

Checks:
  - seq contiguous 1..N
  - every item has invoice/detail/payment
  - delivery files are byte-identical to manifest sources (MD5)
  - every PDF invoice: invoice number and total match manifest
  - verified PDF totals match the manifest amounts of those same PDFs
  - optional workbook: total and row count match
  - optional workbook copy in delivery is identical to the master

Non-PDF invoices (PNG/JPG/OFD) cannot be parsed by pdfplumber. They are reported
as warnings; pass --require-pdf to turn them into errors.

Usage:
    python check_delivery.py --manifest manifest.json --delivery 报销交付_YYYYMMDD \
        --workbook 财务报表.xlsx --buyer-tax-id 91110000XXXXXXXXXX --buyer-name 示例单位
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
try:
    from parse_invoice import parse_pdf
except Exception:
    parse_pdf = None


def md5(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def detail_rank(path: Path) -> int:
    """详情, 详情2, 详情10 -> 1, 2, 10 (lexical sort would put 10 before 2)."""
    m = re.search(r"-详情(\d*)$", path.stem)
    return int(m.group(1) or 1) if m else 0


def to_int(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def read_workbook_totals(path: Path) -> tuple[int, float | None, list[str]]:
    """Return (data_rows, total, warnings). Reads the first worksheet."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[wb.sheetnames[0]]
    warnings: list[str] = []

    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    price_col = None
    for idx, row in enumerate(rows):
        labels = [str(v).strip() if v is not None else "" for v in row]
        if "序号" in labels:
            header_idx = idx
            for col, label in enumerate(labels):
                if "价格" in label or "金额" in label:
                    price_col = col
                    break
            break

    if header_idx is None:
        warnings.append("workbook header row not found; used fallback row counting")
    elif price_col is None:
        warnings.append("workbook price column not found; used last numeric cell in 合计 row")

    data_rows = 0
    total = None
    for row in rows[(header_idx or 0):]:
        values = list(row)
        labels = [str(v).strip() if v is not None else "" for v in values]
        if any("合计" in label for label in labels):
            if price_col is not None and price_col < len(values) and isinstance(values[price_col], (int, float)) \
                    and not isinstance(values[price_col], bool):
                total = float(values[price_col])
            else:
                nums = [float(v) for v in values
                        if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if nums:
                    total = nums[-1]
            continue
        if values and to_int(values[0]) is not None:
            data_rows += 1

    return data_rows, total, warnings


def main() -> int:
    ap = argparse.ArgumentParser(description="Audit reimbursement delivery folder")
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--delivery", required=True)
    ap.add_argument("--root", help="project root (default: manifest directory)")
    ap.add_argument("--workbook", help="master workbook for total comparison")
    ap.add_argument("--buyer-tax-id", default="")
    ap.add_argument("--buyer-name", default="")
    ap.add_argument("--require-pdf", action="store_true",
                    help="fail when a delivery invoice is not PDF")
    args = ap.parse_args()

    manifest_path = Path(args.manifest).resolve()
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    if args.root:
        root = Path(args.root).resolve()
    else:
        root = Path(data.get("root") or manifest_path.parent).resolve()
    delivery = Path(args.delivery).resolve()
    items = sorted([i for i in data.get("items", []) if i.get("status", "complete") == "complete"],
                   key=lambda x: int(x["seq"]))

    errors: list[str] = []
    warnings: list[str] = []

    # 1. sequence
    seqs = [int(i["seq"]) for i in items]
    if seqs != list(range(1, len(items) + 1)):
        errors.append(f"seq not contiguous: {seqs}")

    # 2. files + MD5 + invoice parse
    verified_total = 0.0
    verified_expected = 0.0
    non_pdf_invoices: list[str] = []
    parse_failed: list[str] = []
    n_invoice = n_detail = n_payment = 0

    for it in items:
        seq = int(it["seq"])
        short = it["short"]
        base = f"{seq:02d}-{short}"

        inv_rel = it.get("invoice_file") or ""
        inv_src = root / inv_rel if inv_rel else None
        inv_dst = delivery / f"{base}-发票{inv_src.suffix.lower()}" if inv_src else None
        if inv_dst is None or not inv_dst.is_file():
            errors.append(f"{seq:02d} missing invoice in delivery")
        else:
            n_invoice += 1
            if inv_src is not None and inv_src.is_file():
                if md5(inv_src) != md5(inv_dst):
                    errors.append(f"{seq:02d} invoice MD5 mismatch")
            else:
                warnings.append(f"{seq:02d} master invoice missing: {inv_rel}")

            if inv_dst.suffix.lower() != ".pdf":
                non_pdf_invoices.append(inv_dst.name)
            elif parse_pdf is None:
                parse_failed.append(f"{seq:02d} (pdfplumber unavailable)")
            else:
                try:
                    parsed = parse_pdf(inv_dst, args.buyer_tax_id or None, args.buyer_name or None)
                    verified_total += float(parsed.get("total") or 0)
                    verified_expected += float(it.get("amount") or 0)
                    if parsed.get("invoice_no") != it.get("invoice_no"):
                        errors.append(
                            f"{seq:02d} invoice_no {parsed.get('invoice_no')} != {it.get('invoice_no')}")
                    if abs(float(parsed.get("total") or 0) - float(it.get("amount") or 0)) > 0.005:
                        errors.append(
                            f"{seq:02d} total {parsed.get('total')} != amount {it.get('amount')}")
                    if parsed.get("needs_manual_review"):
                        warnings.append(
                            f"{seq:02d} invoice needs manual review: {parsed.get('checks')}")
                except Exception as exc:
                    parse_failed.append(f"{seq:02d} ({exc})")

        pay_rel = it.get("payment_file") or ""
        pay_src = root / pay_rel if pay_rel else None
        pay_dst = delivery / f"{base}-支付{pay_src.suffix.lower()}" if pay_src else None
        if pay_dst is None or not pay_dst.is_file():
            errors.append(f"{seq:02d} missing payment in delivery")
        else:
            n_payment += 1
            if pay_src is not None and pay_src.is_file():
                if md5(pay_src) != md5(pay_dst):
                    errors.append(f"{seq:02d} payment MD5 mismatch")
            else:
                warnings.append(f"{seq:02d} master payment missing: {pay_rel}")

        order_srcs = [root / p for p in (it.get("order_files") or []) if p]
        details = sorted(delivery.glob(f"{base}-详情*"), key=detail_rank)
        if not details:
            errors.append(f"{seq:02d} missing detail in delivery")
        else:
            n_detail += len(details)
            if len(details) != len(order_srcs):
                errors.append(
                    f"{seq:02d} detail count {len(details)} != order_files {len(order_srcs)}")
            for src, dst in zip(order_srcs, details):
                if src.is_file() and md5(src) != md5(dst):
                    errors.append(f"{seq:02d} detail MD5 mismatch: {dst.name}")

    # 3. totals
    print(f"items={len(items)} invoice={n_invoice} detail={n_detail} payment={n_payment}")
    if non_pdf_invoices:
        warnings.append(
            f"{len(non_pdf_invoices)} non-PDF invoice(s) not machine-verified: "
            + ", ".join(non_pdf_invoices))
        if args.require_pdf:
            errors.append(
                "delivery contains non-PDF invoice(s): " + ", ".join(non_pdf_invoices))
    if parse_failed:
        warnings.append(
            f"{len(parse_failed)} PDF(s) not verified: " + ", ".join(parse_failed))
    if abs(verified_total - verified_expected) > 0.005:
        errors.append(
            f"verified PDF total {verified_total:.2f} != manifest total for verified PDFs {verified_expected:.2f}")
    elif verified_expected:
        print(f"verified invoice total OK: {verified_total:.2f}")

    manifest_total = round(sum(float(i.get("amount") or 0) for i in items), 2)

    # 4. workbook
    if args.workbook:
        wb_path = Path(args.workbook)
        if not wb_path.is_absolute():
            wb_path = root / wb_path
        try:
            data_rows, total_cell, wb_warnings = read_workbook_totals(wb_path)
            warnings.extend(wb_warnings)
            print(f"workbook data_rows={data_rows} total={total_cell}")
            if total_cell is None:
                errors.append("workbook total not found")
            elif abs(total_cell - manifest_total) > 0.005:
                errors.append(f"workbook total {total_cell} != manifest total {manifest_total}")
            if data_rows != len(items):
                errors.append(f"workbook data rows {data_rows} != items {len(items)}")
            copy = delivery / wb_path.name
            if copy.is_file() and md5(copy) != md5(wb_path):
                errors.append("delivery workbook differs from master workbook")
            elif not copy.is_file():
                warnings.append(f"delivery workbook missing: {copy.name}")
        except Exception as exc:
            warnings.append(f"workbook check skipped: {exc}")

    print("\nRESULT:", "FAIL" if errors else "PASS")
    for e in errors:
        print("  ERROR:", e)
    for w in warnings:
        print("  WARN :", w)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
